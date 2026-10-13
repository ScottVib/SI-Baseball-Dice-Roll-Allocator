#!/usr/bin/env python3
"""
Dice-roll allocation for the SI ChartMaker B tabs.

For every B tab and for both grids
    vs LHP : weights C96:C125,  targets row 95  (headers row 94),  entries D96:X125
    vs RHP : weights C137:C166, targets row 136 (headers row 135), entries D137:X166
it assigns each row (dice roll) to exactly ONE of the 21 outcome columns D:X, writing
the row's weight (column C) into that column, so that

  1. each column's total is as close as possible to its target (row 95 / 136):
        minimise  sum_j |column total_j - target_j|
  2. among assignments that are optimal for (1), the totals of related columns
     (1B types, 2B types, hits, outs, GO types, FO types) are as close as
     possible to the sum of their targets:
        minimise  sum_g |group total_g - group target_g|
  3. among those, the single worst column error is as small as possible.

It is solved EXACTLY as a small integer program (30 rows x 21 columns), which takes
a few seconds per grid. Simple greedy / hill-climbing / annealing searches were
tested on the same data and were 5-25% worse, so they are not recommended.

The workbook is only READ (cached values, so save it after Excel has calculated).
Results go to a new workbook with one sheet per B tab, laid out like the grids, so
they can be copied straight over D96:X125 and D137:X166.

Usage:
    pip install openpyxl scipy numpy
    python allocate_dice_rolls.py MyWorkbook.xlsx                 # all tabs B1..B18
    python allocate_dice_rolls.py MyWorkbook.xlsx --tabs B1 B2    # some tabs
    python allocate_dice_rolls.py MyWorkbook.xlsx --out result.xlsx
"""
import argparse
import sys
import numpy as np
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from scipy.optimize import Bounds, LinearConstraint, milp
from scipy.sparse import lil_matrix

HEADERS = ['BB', 'HBP', '(1+)', '(1)', '1', '1*', '1+', '(2)', '2', '2*', '3',
           'HR', 'E', 'SO', 'G', 'G*', 'DP', 'F', '(F)', 'SF', '(SF)']          # D..X

# Groups of related columns, as the sheet itself defines them in rows 129-131.
# Indexes into HEADERS (0 = column D).  Edit here if you want different groups.
GROUPS = {
    '1B types (F:J)':  range(2, 7),
    '2B types (K:M)':  range(7, 10),
    'Hits (F:O)':      range(2, 12),
    'Outs (Q:X)':      range(13, 21),
    'GO types (R:T)':  range(14, 17),
    'FO types (U:X)':  range(17, 21),
}

GRIDS = {                # name: (header row, target row, first row, last row)
    'vs LHP': (94, 95, 96, 125),
    'vs RHP': (135, 136, 137, 166),
}


def solve(weights, targets, time_limit=120):
    """Return (assignment list of column indexes, L1, group deviation, max deviation)."""
    n, m, ng = len(weights), len(targets), len(GROUPS)
    nx = n * m
    off_p, off_n = nx, nx + m
    off_gp, off_gn = nx + 2 * m, nx + 2 * m + ng
    off_mx = off_gn + ng
    nv = off_mx + 1
    integrality = np.zeros(nv)
    integrality[:nx] = 1
    lb = np.zeros(nv)
    ub = np.full(nv, np.inf)
    ub[:nx] = 1

    rows = n + m + ng
    A = lil_matrix((rows, nv))
    lo = np.zeros(rows)
    hi = np.zeros(rows)
    for i in range(n):                                  # every row used exactly once
        for j in range(m):
            A[i, i * m + j] = 1
        lo[i] = hi[i] = 1
    for j in range(m):                                  # column total - target = p - n
        r = n + j
        for i in range(n):
            A[r, i * m + j] = weights[i]
        A[r, off_p + j] = -1
        A[r, off_n + j] = 1
        lo[r] = hi[r] = targets[j]
    for g, cols in enumerate(GROUPS.values()):          # group total - group target = gp - gn
        r = n + m + g
        for j in cols:
            for i in range(n):
                A[r, i * m + j] = weights[i]
        A[r, off_gp + g] = -1
        A[r, off_gn + g] = 1
        lo[r] = hi[r] = sum(targets[j] for j in cols)

    B = lil_matrix((m, nv))                             # |dev_j| <= mx
    for j in range(m):
        B[j, off_p + j] = 1
        B[j, off_n + j] = 1
        B[j, off_mx] = -1
    base = [LinearConstraint(A.tocsr(), lo, hi),
            LinearConstraint(B.tocsr(), np.full(m, -np.inf), np.zeros(m))]

    def run(cost, extra=()):
        res = milp(cost, constraints=base + list(extra), integrality=integrality,
                   bounds=Bounds(lb, ub), options={'time_limit': time_limit})
        if res.x is None:
            raise RuntimeError('solver found no solution: ' + str(res.message))
        return res

    c1 = np.zeros(nv)
    c1[off_p:off_p + 2 * m] = 1
    r1 = run(c1)
    cap1 = LinearConstraint(c1.reshape(1, -1), -np.inf, r1.fun + 1e-6)

    c2 = np.zeros(nv)
    c2[off_gp:off_gn + ng] = 1
    r2 = run(c2, [cap1])
    cap2 = LinearConstraint(c2.reshape(1, -1), -np.inf, r2.fun + 1e-6)

    c3 = np.zeros(nv)
    c3[off_mx] = 1
    r3 = run(c3, [cap1, cap2])

    x = np.round(r3.x[:nx]).reshape(n, m).argmax(axis=1)
    return [int(j) for j in x], r1.fun, r2.fun, r3.fun


def read_grid(ws, grid):
    hdr_row, tgt_row, r1, r2 = GRIDS[grid]
    weights = [ws.cell(r, 3).value for r in range(r1, r2 + 1)]
    labels = [ws.cell(r, 2).value for r in range(r1, r2 + 1)]
    targets = [ws.cell(tgt_row, c).value for c in range(4, 25)]
    if any(not isinstance(v, (int, float)) for v in weights):
        raise ValueError('weights in column C are not all numbers')
    if any(not isinstance(v, (int, float)) for v in targets):
        raise ValueError('targets in row %d are not all numbers (is the workbook calculated and saved?)' % tgt_row)
    if abs(sum(weights) - sum(targets)) > 1e-3:
        raise ValueError('weights total %.2f but targets total %.2f' % (sum(weights), sum(targets)))
    return labels, [float(w) for w in weights], [float(t) for t in targets]


def write_block(ws, top, grid, labels, weights, targets, assign, stats):
    bold = Font(name='Arial', bold=True)
    plain = Font(name='Arial')
    ws.cell(top, 1, grid).font = bold
    ws.cell(top, 2, 'Roll').font = bold
    ws.cell(top, 3, 'Weight').font = bold
    for j, h in enumerate(HEADERS):
        c = ws.cell(top, 4 + j, h)
        c.font = bold
        c.alignment = Alignment(horizontal='center')
    for j, t in enumerate(targets):
        c = ws.cell(top + 1, 4 + j, round(t, 4))
        c.font = plain
        c.number_format = '0.00'
    ws.cell(top + 1, 2, 'Target').font = bold
    for i, (lab, w, col) in enumerate(zip(labels, weights, assign)):
        r = top + 2 + i
        ws.cell(r, 2, lab).font = plain
        ws.cell(r, 3, w).font = plain
        c = ws.cell(r, 4 + col, w)
        c.font = plain
        c.fill = PatternFill('solid', fgColor='DDEBF7')
    r_sum = top + 2 + len(weights)
    ws.cell(r_sum, 2, 'Total').font = bold
    ws.cell(r_sum + 1, 2, 'Over/(Under)').font = bold
    for j in range(len(HEADERS)):
        col_letter = chr(ord('D') + j)
        c = ws.cell(r_sum, 4 + j, '=SUM(%s%d:%s%d)' % (col_letter, top + 2, col_letter, r_sum - 1))
        c.font = plain
        c = ws.cell(r_sum + 1, 4 + j, '=%s%d-%s%d' % (col_letter, r_sum, col_letter, top + 1))
        c.font = plain
        c.number_format = '0.00;(0.00);0.00'
    ws.cell(r_sum + 2, 2, 'Groups').font = bold
    for k, (name, cols) in enumerate(GROUPS.items()):
        S = sum(weights[i] for i, j in enumerate(assign) if j in cols)
        T = sum(targets[j] for j in cols)
        ws.cell(r_sum + 3 + k, 2, name).font = plain
        ws.cell(r_sum + 3 + k, 4, 'target %.2f' % T).font = plain
        ws.cell(r_sum + 3 + k, 5, 'actual %d' % S).font = plain
        ws.cell(r_sum + 3 + k, 6, 'diff %+.2f' % (S - T)).font = plain
    note = ws.cell(r_sum + 3 + len(GROUPS), 2,
                   'Sum of |column errors| = %.3f   group errors = %.3f   worst column error = %.3f' % stats)
    note.font = Font(name='Arial', italic=True)
    return r_sum + 5 + len(GROUPS)


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument('workbook')
    ap.add_argument('--tabs', nargs='*', default=['B%d' % i for i in range(1, 19)])
    ap.add_argument('--out', default='dice_allocations.xlsx')
    ap.add_argument('--time-limit', type=int, default=120, help='seconds per solver stage')
    args = ap.parse_args()

    src = load_workbook(args.workbook, data_only=True)
    out = Workbook()
    out.remove(out.active)
    for tab in args.tabs:
        if tab not in src.sheetnames:
            print('skip %s: no such sheet' % tab)
            continue
        ws_in = src[tab]
        ws_out = out.create_sheet(tab)
        top = 1
        for grid in GRIDS:
            try:
                labels, weights, targets = read_grid(ws_in, grid)
            except ValueError as e:
                print('skip %s %s: %s' % (tab, grid, e))
                continue
            assign, l1, gd, mx = solve(weights, targets, args.time_limit)
            print('%s %s: sum|column error| = %.3f, group error = %.3f, worst column = %.3f' % (tab, grid, l1, gd, mx))
            top = write_block(ws_out, top, grid, labels, weights, targets, assign, (l1, gd, mx))
        ws_out.column_dimensions['B'].width = 14
    out.save(args.out)
    print('written', args.out)


if __name__ == '__main__':
    sys.exit(main())
